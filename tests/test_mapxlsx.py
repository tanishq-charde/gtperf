"""The map workbook's formulas reproduce gtperf.mapgen, and recalculate in a real spreadsheet."""

import io

import libreoffice
import openpyxl
import pytest

from gtperf.app import core
from gtperf.mapgen import CompressorStages, _StackedCompressor
from gtperf.mapxlsx import engine_map_workbook, turbofan_map_workbook

needs_lo = pytest.mark.skipif(not libreoffice.available(), reason="LibreOffice with Python UNO not installed")


@pytest.fixture(scope="module")
def workbook():
    checks = {}
    data = engine_map_workbook(core.design_engine({}), checks)
    return data, checks


def test_formula_values_match_mapgen(workbook):
    _, checks = workbook
    assert set(checks) == {"Compressor", "HP turbine", "Power turbine"}
    for name, c in checks.items():
        assert len(c["points"]) > 100
        for n, w, pr, eta, *_ in c["points"]:
            ref = c["python"].point(n, w)
            assert (pr, eta) == pytest.approx(ref[:2], rel=1e-8), (name, n, w)


def test_sheets_inputs_and_charts(workbook):
    data, _ = workbook
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Read me", "Inputs", "Compressor", "Compressor stages", "HP turbine",
                             "HP turbine stages", "Power turbine", "Power turbine stages"]
    assert {"c_psi", "c_phi", "h_psi", "p_eta", "g_tref"} <= {n.lower() for n in wb.defined_names}
    assert len(wb["Compressor"]._charts) == 2 and len(wb["HP turbine"]._charts) == 2
    formulas = sum(isinstance(c.value, str) and c.value.startswith("=")
                   for ws in wb.worksheets for row in ws.iter_rows() for c in row)
    assert formulas > 50_000


@needs_lo
def test_libreoffice_recalculates_the_same_values_and_follows_inputs(workbook, tmp_path):
    data, checks = workbook
    path = tmp_path / "maps.xlsx"
    path.write_bytes(data)
    cached = openpyxl.load_workbook(path, data_only=True)
    with libreoffice.Calc() as lo:
        doc = lo.open(path)
        for _ in range(2):  # LibreOffice's first hard recalculation of a fresh file can flag false cycles
            doc.calculateAll()
        for name in ("Compressor", "HP turbine", "Power turbine"):
            ws, sheet = cached[name], doc.Sheets.getByName(name)
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, float):
                        cell = sheet.getCellByPosition(c.column - 1, c.row - 1)
                        assert cell.getError() == 0, (name, c.coordinate)
                        assert cell.getValue() == pytest.approx(c.value, rel=1e-8, abs=1e-9), (name, c.coordinate)
        # a different stage loading moves every point the way the Python model does at the same flows
        names = doc.NamedRanges
        addr = names.getByName("c_psi").getReferredCells().getCellAddress()
        doc.Sheets.getByIndex(addr.Sheet).getCellByPosition(addr.Column, addr.Row).setValue(0.41)
        d = core.DESIGN_DEFAULTS
        py = _StackedCompressor(d["pressure_ratio"], d["compressor_polytropic_eff"],
                                CompressorStages(**{**core.compressor_stages(d).__dict__, "loading": 0.41}))
        ws, sheet = cached["Compressor"], doc.Sheets.getByName("Compressor")
        top = next(r for r in range(1, 400) if ws.cell(r, 1).value == "Map points (R-line 1 = surge, 3 = choke end)")
        n_checked = 0
        for r in range(top + 3, top + 3 + len(checks["Compressor"]["points"]), 7):
            n, w = (sheet.getCellByPosition(c, r - 1).getValue() for c in (0, 2))
            ref = py.point(n, w)
            if ref is not None:
                assert sheet.getCellByPosition(4, r - 1).getValue() == pytest.approx(ref[0], rel=1e-7)
                n_checked += 1
        assert n_checked > 20
        doc.close(True)


def test_turbofan_workbook_has_every_component():
    checks = {}
    data = turbofan_map_workbook(core.turbofan_engine(), None, checks)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    assert wb.sheetnames == ["Read me", "Inputs", "Fan", "Fan stages", "Booster", "Booster stages", "HPC",
                             "HPC stages", "HPT", "HPT stages", "LPT", "LPT stages"]
    for name, c in checks.items():
        for n, w, pr, eta, *_ in c["points"]:
            assert (pr, eta) == pytest.approx(c["python"].point(n, w)[:2], rel=1e-8), (name, n, w)
