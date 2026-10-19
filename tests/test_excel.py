import io
import tomllib
from contextlib import redirect_stdout
from pathlib import Path

import libreoffice
import pytest
from openpyxl import load_workbook
from test_cli_and_study import DATASHEET

from gtperf.cli import engine_from_datasheet, main
from gtperf.excel import DEFAULT_CASES, OUTPUTS, Case, build_workbook, correction_tables, evaluate, full_model

BAS = Path(__file__).resolve().parents[1] / "excel" / "gtperf.bas"
needs_lo = pytest.mark.skipif(not libreoffice.available(), reason="LibreOffice with Python UNO not installed")

OFF_GRID = Case("Off grid", 37.0, 800.0, 0.40, 12.0, 30.0, 1.5, 0.8)


@pytest.fixture(scope="module")
def engine():
    return engine_from_datasheet(DATASHEET)[2]


@pytest.fixture(scope="module")
def tables(engine):
    return correction_tables(engine)


@pytest.fixture(scope="module")
def workbook(engine, tables, tmp_path_factory):
    path = tmp_path_factory.mktemp("xl") / "site_tool.xlsx"
    build_workbook(engine, path, name="test engine", tables=tables, datasheet_path=DATASHEET)
    return path


def rows_by_label(path):
    ws = load_workbook(path)["Site tool"]
    out = {}
    for r in range(1, ws.max_row + 1):
        out.setdefault(ws.cell(r, 2).value, r)
    return out


def test_tables_reproduce_iso_and_show_speed_limit(tables):
    i = tables["temperature_C"].index(15)
    assert tables["t_power"][i] == pytest.approx(1.0, abs=1e-6)
    assert tables["control"][0] == "gas generator corrected speed" and tables["control"][-1] == "temperature"
    assert all(k < 0 for k in tables["inlet_power"]) and all(k > 0 for k in tables["inlet_heat_rate"])


@pytest.mark.parametrize("case", [*DEFAULT_CASES, OFF_GRID], ids=lambda c: c.name)
def test_correction_method_close_to_full_model(engine, tables, case):
    e, f = evaluate(tables, case), full_model(engine, case)
    for q in ("power", "heat_rate", "exhaust_flow"):
        assert e[q] == pytest.approx(f[q], rel=0.003), q
    assert e["exhaust_temp"] == pytest.approx(f["exhaust_temp"], abs=0.5)


def test_workbook_structure(workbook):
    wb = load_workbook(workbook)
    assert wb.sheetnames == ["Site tool", "Tables", "Check", "About"]
    for name in ("T_GRID", "T_POWER", "RH_HEAT_RATE", "EXHAUST_EXHAUST_TEMP", "ISO_POWER_MW", "DATASHEET_PATH"):
        assert name in wb.defined_names
    rows = rows_by_label(workbook)
    assert wb["Site tool"].cell(rows["Case name"], 4).value == "ISO"
    assert str(wb["Site tool"].cell(rows["Electrical output"], 4).value).startswith("=ISO_POWER_MW*")


def test_cli_writes_workbook(tmp_path):
    out = tmp_path / "tool.xlsx"
    with redirect_stdout(io.StringIO()):
        main(["excel", str(DATASHEET), str(out)])
    assert load_workbook(out)["About"]["A2"].value.startswith("Engine: SGT-400")


@needs_lo
def test_formulas_and_vba_in_libreoffice(workbook, tables, tmp_path):
    """Recalculate the workbook in LibreOffice and run the VBA module in its VBA mode.

    Both must agree with `evaluate`, including after a case is edited, and the VBA
    site-file writer and CSV reader must round-trip through the real CLI.
    """
    rows = rows_by_label(workbook)
    cols = "DEFGH"
    result_rows = {"power": rows["Electrical output"], "heat_rate": rows["Heat rate (LHV)"],
                   "exhaust_flow": rows["Exhaust flow"], "exhaust_temp": rows["Exhaust temperature"]}
    inputs = [("temperature_C", "Ambient temperature", 1), ("altitude_m", "Site altitude", 1),
              ("rh", "Relative humidity", 100), ("inlet_loss_mbar", "Inlet pressure loss", 1),
              ("exhaust_loss_mbar", "Exhaust pressure loss", 1),
              ("degradation_power_pct", "Degradation, output", 1),
              ("degradation_heat_rate_pct", "Degradation, heat rate", 1)]
    with libreoffice.Calc() as lo:
        doc = lo.open(workbook, macros=True)
        sheet = doc.Sheets.getByName("Site tool")
        for field, label, scale in inputs:  # replace the last case with an off-grid one
            sheet.getCellRangeByName(f"H{rows[label]}").Value = getattr(OFF_GRID, field) * scale
        lo.add_vba_module(doc, "gtperf", BAS.read_text())
        a = (OFF_GRID.temperature_C, OFF_GRID.altitude_m, OFF_GRID.rh * 100, OFF_GRID.inlet_loss_mbar,
             OFF_GRID.exhaust_loss_mbar)
        args = ";".join(str(v) for v in a)
        deg = f"{OFF_GRID.degradation_power_pct};{OFF_GRID.degradation_heat_rate_pct}"
        vba = {q: f'=GT_SITE("{q}";{args};{deg})' for q in OUTPUTS}
        vba["toml"] = f'=BUILDSITETOML("Off grid";{args})'
        vba["isa"] = "=GT_ISA_PRESSURE(800)"
        for i, f in enumerate(vba.values()):
            libreoffice.set_formula(doc, "About", f"C{50 + i}", f)
        doc.calculateAll()
        sheet_values = {(q, j): libreoffice.cell(doc, "Site tool", f"{cols[j]}{r}")
                        for q, r in result_rows.items() for j in range(5)}
        vba_values = {k: libreoffice.cell(doc, "About", f"C{50 + i}") for i, k in enumerate(vba)}

        site_file = tmp_path / "site.toml"
        site_file.write_text(vba_values["toml"])
        buf = io.StringIO()
        with redirect_stdout(buf):
            main(["site", str(DATASHEET), str(site_file), "--csv"])
        csv = buf.getvalue().replace("\r", "").replace('"', '""').replace("\n", '"&CHAR(10)&"')
        libreoffice.set_formula(doc, "About", "C70", f'=PARSESITECSV("{csv}";"Electrical output")')
        doc.calculateAll()
        parsed_power = libreoffice.cell(doc, "About", "C70")
        doc.close(True)

    for j, case in enumerate([*DEFAULT_CASES[:4], OFF_GRID]):
        expected = evaluate(tables, case)
        for q in OUTPUTS:
            assert sheet_values[q, j] == pytest.approx(expected[q], rel=1e-9), (case.name, q)
    expected = evaluate(tables, OFF_GRID)
    for q in OUTPUTS:
        assert vba_values[q] == pytest.approx(expected[q], rel=1e-9), q
    assert vba_values["isa"] == pytest.approx(expected["pressure_mbar"], rel=1e-12)
    site = tomllib.loads(vba_values["toml"])
    assert site == {"name": "Off grid", "temperature_C": 37, "altitude_m": 800, "rh": 0.4, "inlet_loss_mbar": 12,
                    "exhaust_loss_mbar": 30}
    full = full_model(engine_from_datasheet(DATASHEET)[2], Case("x", *a[:2], 0.4, *a[3:]))
    assert parsed_power == pytest.approx(full["power"], abs=0.006)  # CSV is printed to 2 decimals
