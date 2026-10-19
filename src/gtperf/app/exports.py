"""Download formats for GUI results: SVG figures, an Excel workbook, a PowerPoint deck and a PDF report.

Each exporter takes a list of `core.Section` objects and returns bytes, so the GUI can offer them as downloads
and tests can check them without a browser.
"""

from __future__ import annotations

import datetime as _dt
import io
import re
import zipfile
from pathlib import Path

from .core import INK, INK2, NAVY

ASSETS = Path(__file__).resolve().parent / "assets"
LOGO = ASSETS / "logo.png"
TOOL = "GTPerf"
STATUS = "beta"


def _slug(text):
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_").lower()


def _fmt(v):
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, float):
        if v != v:  # NaN
            return ""
        a = abs(v)
        if a < 1e-9:
            return "0"
        if a >= 1000:
            return f"{v:,.0f}"
        if a >= 10:
            return f"{v:.2f}"
        return f"{v:.4g}"
    return "" if v is None else str(v)


def figure_bytes(fig, fmt="svg", dpi=200):
    buf = io.BytesIO()
    fig.savefig(buf, format=fmt, dpi=dpi, facecolor=fig.get_facecolor())
    return buf.getvalue()


def figure_svg(fig):
    return figure_bytes(fig, "svg")


def figures_zip(sections):
    """Every figure as SVG, in one zip, one folder per section."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for s in sections:
            for name, fig in s.figures.items():
                z.writestr(f"{_slug(s.title)}/{_slug(name)}.svg", figure_svg(fig))
    return buf.getvalue()


def tables_csv_zip(sections):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for s in sections:
            for name, df in s.tables.items():
                z.writestr(f"{_slug(s.title)}/{_slug(name)}.csv", df.to_csv(index=False))
    return buf.getvalue()


# --- Excel -------------------------------------------------------------------------------------------


def to_xlsx(sections):
    """One summary sheet (inputs and key results of every section), then one sheet per table."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=NAVY.lstrip("#"))
    ws.append([f"{TOOL} results ({STATUS})", f"generated {_dt.datetime.now():%Y-%m-%d %H:%M}"])
    ws["A1"].font = Font(bold=True, size=14, color=NAVY.lstrip("#"))
    for s in sections:
        ws.append([])
        ws.append([s.title])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12, color=NAVY.lstrip("#"))
        ws.append(["Input", "Value"])
        for c in ws[ws.max_row]:
            c.font, c.fill = head, fill
        for k, v in s.inputs.items():
            ws.append([k, v if isinstance(v, (int, float)) else _fmt(v)])
        ws.append(["Key result", "Value"])
        for c in ws[ws.max_row]:
            c.font, c.fill = head, fill
        for k, v in s.key_results:
            ws.append([k, v])
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 48
    used = {"Summary"}
    for s in sections:
        for name, df in s.tables.items():
            title = re.sub(r"[\[\]\*\?/\\:]", "", f"{s.title[:12]} - {name}")[:31]
            base, i = title, 2
            while title in used:
                title = f"{base[:28]}~{i}"
                i += 1
            used.add(title)
            sh = wb.create_sheet(title)
            sh.append(list(df.columns))
            for c in sh[1]:
                c.font, c.fill = head, fill
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for row in df.itertuples(index=False):
                sh.append(
                    [
                        None if (isinstance(v, float) and v != v) else (v.item() if hasattr(v, "item") else v)
                        for v in row
                    ]
                )
            for j, col in enumerate(df.columns, 1):
                sh.column_dimensions[get_column_letter(j)].width = max(10, min(40, len(str(col)) + 2))
            sh.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- PowerPoint --------------------------------------------------------------------------------------


def to_pptx(sections, title="Gas turbine performance results"):
    """A 16:9 deck: cover, then per section a key-results slide, one slide per figure and the first table."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    navy = RGBColor.from_string(NAVY.lstrip("#"))
    ink2 = RGBColor.from_string(INK2.lstrip("#"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, s, size=18, bold=False, color=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s
        p.font.size, p.font.bold = Pt(size), bold
        p.font.color.rgb = color or RGBColor.from_string(INK.lstrip("#"))
        return tf

    def header(slide, s, kicker):
        text(slide, 0.6, 0.35, 12, 0.4, kicker.upper(), 12, True, navy)
        text(slide, 0.6, 0.7, 12, 0.8, s, 28, True)
        if LOGO.exists():
            slide.shapes.add_picture(str(LOGO), Inches(11.4), Inches(0.3), height=Inches(0.55))
        bar = slide.shapes.add_shape(1, Inches(0), Inches(7.3), prs.slide_width, Inches(0.2))
        bar.fill.solid()
        bar.fill.fore_color.rgb = navy
        bar.line.fill.background()

    cover = prs.slides.add_slide(blank)
    if LOGO.exists():
        cover.shapes.add_picture(str(LOGO), Inches(0.8), Inches(1.4), height=Inches(1.6))
    text(cover, 0.8, 3.4, 11.5, 1, title, 36, True)
    text(cover, 0.8, 4.4, 11.5, 0.6, f"{TOOL} {STATUS} · generated {_dt.date.today():%d %B %Y}", 16, False, ink2)
    text(cover, 0.8, 5.0, 11.5, 1.2, "Sections: " + ", ".join(s.title for s in sections), 14, False, ink2)

    for s in sections:
        sl = prs.slides.add_slide(blank)
        header(sl, s.title, "Key results")
        tf = text(sl, 0.6, 1.7, 6.2, 5.2, "", 16)
        for i, (k, v) in enumerate(s.key_results):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r1 = p.add_run()
            r1.text = f"{k}: "
            r1.font.size, r1.font.bold = Pt(16), True
            r2 = p.add_run()
            r2.text = str(v)
            r2.font.size = Pt(16)
            p.space_after = Pt(8)
        tf2 = text(sl, 7.2, 1.7, 5.6, 5.2, "Inputs", 14, True, navy)
        for k, v in list(s.inputs.items())[:14]:
            p = tf2.add_paragraph()
            p.text = f"{k}: {_fmt(v)}"
            p.font.size = Pt(12)
            p.font.color.rgb = ink2
        for name, fig in s.figures.items():
            sl = prs.slides.add_slide(blank)
            header(sl, name, s.title)
            png = io.BytesIO(figure_bytes(fig, "png", 200))
            w_in, h_in = fig.get_size_inches()
            max_w, max_h = 12.1, 5.5
            scale = min(max_w / w_in, max_h / h_in)
            pic_w, pic_h = w_in * scale, h_in * scale
            sl.shapes.add_picture(png, Inches(0.6 + (max_w - pic_w) / 2), Inches(1.6), Inches(pic_w), Inches(pic_h))
        if s.tables:
            name, df = next(iter(s.tables.items()))
            df = df.head(12)
            cols = list(df.columns)[:8]
            sl = prs.slides.add_slide(blank)
            header(sl, name, s.title)
            shape = sl.shapes.add_table(
                len(df) + 1, len(cols), Inches(0.6), Inches(1.6), Inches(12.1), Emu(int(Inches(0.38)) * (len(df) + 1))
            )
            tbl = shape.table
            for j, c in enumerate(cols):
                cell = tbl.cell(0, j)
                cell.text = str(c)
                cell.text_frame.paragraphs[0].font.size = Pt(11)
                cell.text_frame.paragraphs[0].font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = navy
            for i, row in enumerate(df[cols].itertuples(index=False), 1):
                for j, v in enumerate(row):
                    cell = tbl.cell(i, j)
                    cell.text = _fmt(v)
                    cell.text_frame.paragraphs[0].font.size = Pt(10)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# --- PDF -----------------------------------------------------------------------------------------------


def _font_path(bold=False):
    import matplotlib

    name = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    return Path(matplotlib.get_data_path()) / "fonts" / "ttf" / name


def to_pdf(sections, title="Gas turbine performance results"):
    """A4 report: cover, then per section the inputs, key results, notes, figures and tables."""
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    navy = tuple(int(NAVY[i : i + 2], 16) for i in (1, 3, 5))
    ink2 = tuple(int(INK2[i : i + 2], 16) for i in (1, 3, 5))

    class Report(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font("body", size=8)
            self.set_text_color(*ink2)
            self.cell(0, 6, f"{TOOL} {STATUS} · {title} · page {self.page_no()}", align="C")

    pdf = Report(format="A4")
    pdf.add_font("body", "", str(_font_path()))
    pdf.add_font("body", "B", str(_font_path(True)))
    pdf.set_auto_page_break(True, margin=16)
    pdf.set_margins(16, 16, 16)
    width = pdf.w - 32

    pdf.add_page()
    if LOGO.exists():
        pdf.image(str(LOGO), x=16, y=40, h=34)
    pdf.set_y(90)
    pdf.set_font("body", "B", 24)
    pdf.set_text_color(*navy)
    pdf.multi_cell(width, 11, title)
    pdf.set_font("body", "", 12)
    pdf.set_text_color(*ink2)
    pdf.multi_cell(width, 7, f"{TOOL} {STATUS} · generated {_dt.datetime.now():%d %B %Y, %H:%M}")
    pdf.ln(6)
    for i, s in enumerate(sections, 1):
        pdf.multi_cell(width, 7, f"{i}. {s.title}")

    def heading(text_, size=16):
        pdf.set_font("body", "B", size)
        pdf.set_text_color(*navy)
        pdf.multi_cell(width, size * 0.55, text_)
        pdf.ln(2)

    def table(rows, header, col_widths=None, max_rows=40):
        pdf.set_font("body", "", 8)
        pdf.set_text_color(0, 0, 0)
        style = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=navy)
        with pdf.table(
            width=width,
            col_widths=col_widths,
            headings_style=style,
            line_height=5,
            text_align="LEFT",
            borders_layout="HORIZONTAL_LINES",
        ) as t:
            t.row([str(h) for h in header])
            for r in rows[:max_rows]:
                t.row([_fmt(v) for v in r])
        if len(rows) > max_rows:
            pdf.set_font("body", "", 8)
            pdf.set_text_color(*ink2)
            pdf.multi_cell(width, 5, f"First {max_rows} of {len(rows)} rows; the Excel download has them all.")
        pdf.ln(4)

    for i, s in enumerate(sections, 1):
        pdf.add_page()
        heading(f"{i}. {s.title}", 18)
        heading("Key results", 12)
        table([(k, v) for k, v in s.key_results], ["Result", "Value"], (1, 1.4))
        heading("Inputs", 12)
        table([(k, v) for k, v in s.inputs.items()], ["Input", "Value"], (1, 1.4))
        if s.notes:
            pdf.set_font("body", "", 9)
            pdf.set_text_color(*ink2)
            for n in s.notes:
                pdf.multi_cell(width, 5, f"• {n}")
            pdf.ln(3)
        for name, fig in s.figures.items():
            w_in, h_in = fig.get_size_inches()
            h_mm = width * h_in / w_in
            if pdf.get_y() + h_mm + 10 > pdf.h - 16:
                pdf.add_page()
            heading(name, 11)
            pdf.image(io.BytesIO(figure_bytes(fig, "png", 200)), x=16, w=width)
            pdf.ln(4)
        for name, df in s.tables.items():
            if len(df.columns) > 9:
                df = df.iloc[:, :9]
            if pdf.get_y() > pdf.h - 50:
                pdf.add_page()
            heading(name, 11)
            table([tuple(r) for r in df.itertuples(index=False)], list(df.columns), max_rows=25)
    return bytes(pdf.output())
