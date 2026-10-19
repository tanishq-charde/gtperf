"""Excel site performance tool generated from a calibrated engine model.

The workbook is the kind of correction-curve sheet a sales or proposals team uses:
a non-specialist types in site conditions and gets output, heat rate and exhaust
conditions from tabulated corrections, with no Python and no macros needed. The
tables are produced here from the full off-design model, so they can be rebuilt
for any calibrated engine with `gtperf excel`.

Corrections are applied in the same order as `site_performance`:

    temperature curve  ->  altitude (pressure ratio delta)  ->  humidity  ->
    inlet loss  ->  exhaust loss  ->  degradation

Humidity and the two losses are linear sensitivities tabulated against ambient
temperature; a pressure loss is divided by delta because its effect scales with
the loss as a fraction of ambient pressure. `evaluate` mirrors the spreadsheet
formulas exactly, so the formula approximation can be checked against the full
model without opening Excel.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass

import numpy as np

from .cycle import ISO, Ambient
from .site import Site, isa_pressure, site_performance

T_GRID = tuple(range(-20, 51, 5))  # degC
RH_REF = ISO.rh
INLET_STEP, EXHAUST_STEP = 10.0, 25.0  # mbar used for the loss sensitivities
OUTPUTS = ("power", "heat_rate", "exhaust_flow", "exhaust_temp")


@dataclass
class Case:
    name: str
    temperature_C: float = 15.0
    altitude_m: float = 0.0
    rh: float = 0.60
    inlet_loss_mbar: float = 0.0
    exhaust_loss_mbar: float = 0.0
    degradation_power_pct: float = 0.0
    degradation_heat_rate_pct: float = 0.0


DEFAULT_CASES = (
    Case("ISO"),
    Case("Temperate, sea level", 10.0, 30.0, 0.80, 8.0, 15.0),
    Case("Hot desert", 45.0, 300.0, 0.20, 10.0, 25.0),
    Case("Tropical coast", 32.0, 10.0, 0.85, 8.0, 20.0),
    Case("Cold, elevated", -10.0, 1500.0, 0.60, 8.0, 15.0, 2.0, 1.0),
)


def _quantities(r):
    return np.array([r.electrical_power, r.heat_rate, r.exhaust.w, r.exhaust.t])


def correction_tables(engine, temperatures=T_GRID):
    """Tabulate the corrections against ambient temperature at sea level.

    Returns a dict of lists. Ratios for power, heat rate and exhaust flow;
    exhaust temperature changes are in kelvin. Humidity sensitivities are per
    unit relative humidity, loss sensitivities per mbar at sea level.
    """
    iso = _quantities(engine.run(ISO, inlet_loss=0.0, exhaust_loss=0.0))
    tab = {"temperature_C": list(temperatures)}
    for prefix in ("t", "rh", "inlet", "exhaust"):
        for q in OUTPUTS:
            tab[f"{prefix}_{q}"] = []
    tab["control"] = []
    guess = None

    def run(t, rh, inlet=0.0, exhaust=0.0):
        nonlocal guess
        r = engine.run(Ambient(t + 273.15, ISO.p, rh), inlet_loss=inlet * 100, exhaust_loss=exhaust * 100,
                       guess=guess)
        guess = r.info.get("solution")
        return r

    for t in temperatures:
        base = run(t, RH_REF)
        b = _quantities(base)
        dry, wet = _quantities(run(t, 0.0)), _quantities(run(t, 1.0))
        inl = _quantities(run(t, RH_REF, inlet=INLET_STEP))
        exh = _quantities(run(t, RH_REF, exhaust=EXHAUST_STEP))
        rel = [b[0] / iso[0], b[1] / iso[1], b[2] / iso[2], b[3] - iso[3]]
        for q, v in zip(OUTPUTS, rel):
            tab[f"t_{q}"].append(float(v))
        for prefix, other, scale in (("rh", wet, 1.0), ("inlet", inl, INLET_STEP), ("exhaust", exh, EXHAUST_STEP)):
            low = dry if prefix == "rh" else b
            sens = [(other[k] - low[k]) / b[k] / scale for k in range(3)] + [(other[3] - low[3]) / scale]
            for q, v in zip(OUTPUTS, sens):
                tab[f"{prefix}_{q}"].append(float(v))
        tab["control"].append(base.info["control"])
    tab["iso"] = {"power_MW": iso[0] / 1e6, "heat_rate": iso[1], "exhaust_flow": iso[2],
                  "exhaust_temp_C": iso[3] - 273.15}
    return tab


def _interp(x, xs, ys):
    """Linear interpolation that extrapolates from the end segments, like the workbook's FORECAST formula."""
    xs, ys = np.asarray(xs, float), np.asarray(ys, float)
    i = int(np.clip(np.searchsorted(xs, x, side="right") - 1, 0, len(xs) - 2))
    return ys[i] + (ys[i + 1] - ys[i]) * (x - xs[i]) / (xs[i + 1] - xs[i])


def evaluate(tab, case):
    """Site performance from the tables, exactly as the workbook formulas compute it."""
    T = tab["temperature_C"]
    g = {k: _interp(case.temperature_C, T, v) for k, v in tab.items() if isinstance(v, list) and k != "control"
         and k != "temperature_C"}
    delta = isa_pressure(case.altitude_m) / ISO.p
    drh = case.rh - RH_REF
    din, dex = case.inlet_loss_mbar / delta, case.exhaust_loss_mbar / delta
    iso = tab["iso"]
    out = {}
    for q, base, alt in (("power", iso["power_MW"], delta), ("heat_rate", iso["heat_rate"], 1.0),
                         ("exhaust_flow", iso["exhaust_flow"], delta)):
        f = g[f"t_{q}"] * alt * (1 + g[f"rh_{q}"] * drh) * (1 + g[f"inlet_{q}"] * din) * (1 + g[f"exhaust_{q}"] * dex)
        out[q] = base * f
    out["power"] *= 1 - case.degradation_power_pct / 100
    out["heat_rate"] *= 1 + case.degradation_heat_rate_pct / 100
    out["exhaust_temp"] = (iso["exhaust_temp_C"] + g["t_exhaust_temp"] + g["rh_exhaust_temp"] * drh
                           + g["inlet_exhaust_temp"] * din + g["exhaust_exhaust_temp"] * dex)
    out["pressure_mbar"] = delta * ISO.p / 100
    return out


def full_model(engine, case):
    """The same case through the full model (degradation applied as the same simple factors)."""
    site = Site(case.name, case.altitude_m, case.temperature_C, case.rh, case.inlet_loss_mbar, case.exhaust_loss_mbar)
    r = site_performance(engine, site).result
    return {"power": r.electrical_power / 1e6 * (1 - case.degradation_power_pct / 100),
            "heat_rate": r.heat_rate * (1 + case.degradation_heat_rate_pct / 100),
            "exhaust_flow": r.exhaust.w, "exhaust_temp": r.exhaust.t - 273.15}


# --- workbook ------------------------------------------------------------------

INK, INK2, GRID, BLUE_LIGHT, BLUE_DARK, SURFACE = "0B0B0B", "52514E", "E4E3DF", "DCEAFB", "184F95", "FCFCFB"

# (label, unit, Case field, workbook value -> Case value, default format, validation min, max)
INPUT_ROWS = (
    ("Ambient temperature", "°C", "temperature_C", 1, "0.0", -20, 50),
    ("Site altitude", "m", "altitude_m", 1, "0", 0, 3000),
    ("Relative humidity", "%", "rh", 100, "0", 0, 100),
    ("Inlet pressure loss", "mbar", "inlet_loss_mbar", 1, "0.0", 0, 50),
    ("Exhaust pressure loss", "mbar", "exhaust_loss_mbar", 1, "0.0", 0, 80),
    ("Degradation, output", "%", "degradation_power_pct", 1, "0.0", 0, 15),
    ("Degradation, heat rate", "%", "degradation_heat_rate_pct", 1, "0.0", 0, 15),
)


def _formula_interp(x, name):
    """Excel formula for linear interpolation of the named table column against T_GRID."""
    i = f"MIN(IFERROR(MATCH({x},T_GRID,1),1),ROWS(T_GRID)-1)-1"
    return f"FORECAST({x},OFFSET({name},{i},0,2,1),OFFSET(T_GRID,{i},0,2,1))"


def build_workbook(engine, path, name="engine", cases=DEFAULT_CASES, tables=None, source="", datasheet_path="",
                   command="gtperf"):
    """Write the site performance workbook. Returns the tables used."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, PatternFill, Side
    from openpyxl.styles import Font as _Font
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    def Font(**kw):  # noqa: N802 - same call signature as openpyxl's Font
        return _Font(name=kw.pop("name", "Calibri"), **kw)

    tab = tables or correction_tables(engine)
    wb = Workbook()
    ws = wb.active
    ws.title = "Site tool"
    tables_ws = wb.create_sheet("Tables")
    check_ws = wb.create_sheet("Check")
    about_ws = wb.create_sheet("About")

    bold, head = Font(bold=True, color=INK), Font(bold=True, color="FFFFFF")
    muted = Font(color=INK2, italic=True)
    head_fill = PatternFill("solid", fgColor=BLUE_DARK)
    input_fill = PatternFill("solid", fgColor=BLUE_LIGHT)
    thin = Side(style="thin", color=GRID)
    box = Border(top=thin, bottom=thin, left=thin, right=thin)

    # Tables sheet with named columns
    cols = [k for k, v in tab.items() if isinstance(v, list)]
    for j, k in enumerate(cols, start=1):
        c = tables_ws.cell(1, j, k)
        c.font, c.fill = head, head_fill
        for i, v in enumerate(tab[k], start=2):
            tables_ws.cell(i, j, v).number_format = "0.000000" if isinstance(v, float) else "General"
        tables_ws.column_dimensions[get_column_letter(j)].width = 16
        if k != "control":
            ref = f"Tables!${get_column_letter(j)}$2:${get_column_letter(j)}${len(tab[k]) + 1}"
            label = "T_GRID" if k == "temperature_C" else k.upper()
            wb.defined_names[label] = DefinedName(label, attr_text=ref)
    n = len(tab["temperature_C"]) + 3
    tables_ws.cell(n, 1, "Ratios to ISO for t_*; exhaust temperature columns are kelvin. rh_* per unit RH about "
                         f"{RH_REF:.0%}, inlet_* per mbar of inlet loss, exhaust_* per mbar of exhaust loss, all at "
                         "sea level.").font = muted
    iso_row = n + 2
    for j, (k, v) in enumerate(tab["iso"].items(), start=1):
        tables_ws.cell(iso_row, j, k).font = bold
        tables_ws.cell(iso_row + 1, j, v)
        letter = get_column_letter(j)
        wb.defined_names[f"ISO_{k.upper()}"] = DefinedName(f"ISO_{k.upper()}",
                                                          attr_text=f"Tables!${letter}${iso_row + 1}")
    tables_ws.freeze_panes = "B2"

    # Site tool sheet
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 9
    ncase = len(cases)
    first = 4  # column D
    for j in range(ncase):
        ws.column_dimensions[get_column_letter(first + j)].width = 17
    ws["B1"] = f"Site performance: {name}"
    ws["B1"].font = Font(bold=True, size=14, color=INK)
    ws["B2"] = ("Blue cells are inputs. Results use correction curves from the calibrated gtperf model; "
                "see the Check sheet for their error against the full model.")
    ws["B2"].font = muted

    row = 4

    def header(text):
        nonlocal row
        ws.cell(row, 2, text).font = head
        for j in range(2, first + ncase):
            ws.cell(row, j).fill = head_fill
        row += 1

    def put(label, unit, values, fmt, fill=None):
        nonlocal row
        ws.cell(row, 2, label).font = Font(color=INK)
        ws.cell(row, 3, unit).font = Font(color=INK2)
        for j, v in enumerate(values):
            c = ws.cell(row, first + j, v)
            c.number_format, c.border = fmt, box
            c.alignment = Alignment(horizontal="right")
            if fill:
                c.fill = fill
        row += 1
        return row - 1

    def ref(r, j):
        return f"{get_column_letter(first + j)}{r}"

    header("Case")
    case_row = put("Case name", "", [c.name for c in cases], "@", input_fill)
    for j in range(ncase):
        ws.cell(case_row, first + j).alignment = Alignment(horizontal="center", wrap_text=True)
    rows = {}
    header("Site conditions")
    for label, unit, field_, scale, fmt, lo, hi in INPUT_ROWS:
        r = put(label, unit, [getattr(c, field_) * scale for c in cases], fmt, input_fill)
        rows[field_] = r
        dv = DataValidation(type="decimal", operator="between", formula1=str(lo), formula2=str(hi),
                            showErrorMessage=True, errorTitle="Out of range",
                            error=f"{label} must be between {lo} and {hi} {unit}.")
        dv.add(f"{ref(r, 0)}:{ref(r, ncase - 1)}")
        ws.add_data_validation(dv)

    def x(field_, j):
        return ref(rows[field_], j)

    header("Derived")
    p_row = put("Ambient pressure (ISA)", "mbar",
                [f"=1013.25*(1-2.25577E-5*{x('altitude_m', j)})^5.25588" for j in range(ncase)], "0.0")
    d_row = put("Pressure ratio to ISO, delta", "-", [f"={ref(p_row, j)}/1013.25" for j in range(ncase)], "0.0000")

    def g(name_, j):
        return _formula_interp(x("temperature_C", j), name_.upper())

    header("Correction factors (site = ISO x product)")
    fac = {}
    for q, label in (("power", "output"), ("heat_rate", "heat rate"), ("exhaust_flow", "exhaust flow")):
        alt = (lambda j: ref(d_row, j)) if q != "heat_rate" else (lambda j: "1")
        fac[q] = [
            put(f"Temperature, {label}", "-", [f"={g('t_' + q, j)}" for j in range(ncase)], "0.0000"),
            put(f"Altitude, {label}", "-", [f"={alt(j)}" for j in range(ncase)], "0.0000"),
            put(f"Humidity, {label}", "-",
                [f"=1+{g('rh_' + q, j)}*({x('rh', j)}/100-{RH_REF})" for j in range(ncase)], "0.0000"),
            put(f"Inlet loss, {label}", "-",
                [f"=1+{g('inlet_' + q, j)}*{x('inlet_loss_mbar', j)}/{ref(d_row, j)}" for j in range(ncase)],
                "0.0000"),
            put(f"Exhaust loss, {label}", "-",
                [f"=1+{g('exhaust_' + q, j)}*{x('exhaust_loss_mbar', j)}/{ref(d_row, j)}" for j in range(ncase)],
                "0.0000"),
        ]
    fac["power"].append(put("Degradation factor, output", "-",
                            [f"=1-{x('degradation_power_pct', j)}/100" for j in range(ncase)], "0.0000"))
    fac["heat_rate"].append(put("Degradation factor, heat rate", "-",
                                [f"=1+{x('degradation_heat_rate_pct', j)}/100" for j in range(ncase)], "0.0000"))

    def product(q, j):
        return "*".join(ref(r, j) for r in fac[q])

    header("Site performance")
    res = {}
    res["power"] = put("Electrical output", "MW", [f"=ISO_POWER_MW*{product('power', j)}" for j in range(ncase)],
                       "0.00")
    res["heat_rate"] = put("Heat rate (LHV)", "kJ/kWh",
                           [f"=ISO_HEAT_RATE*{product('heat_rate', j)}" for j in range(ncase)], "#,##0")
    put("Thermal efficiency (LHV)", "%", [f"=360000/{ref(res['heat_rate'], j)}" for j in range(ncase)], "0.00")
    put("Fuel heat input (LHV)", "MW",
        [f"={ref(res['power'], j)}*{ref(res['heat_rate'], j)}/3600" for j in range(ncase)], "0.00")
    res["exhaust_flow"] = put("Exhaust flow", "kg/s",
                              [f"=ISO_EXHAUST_FLOW*{product('exhaust_flow', j)}" for j in range(ncase)], "0.00")
    res["exhaust_temp"] = put(
        "Exhaust temperature", "°C",
        [f"=ISO_EXHAUST_TEMP_C+{g('t_exhaust_temp', j)}+{g('rh_exhaust_temp', j)}*({x('rh', j)}/100-{RH_REF})"
         f"+{g('inlet_exhaust_temp', j)}*{x('inlet_loss_mbar', j)}/{ref(d_row, j)}"
         f"+{g('exhaust_exhaust_temp', j)}*{x('exhaust_loss_mbar', j)}/{ref(d_row, j)}" for j in range(ncase)],
        "0.0")
    row += 1
    ws.cell(row, 2, "Fuel composition changes and part load need the full model (gtperf site, or the "
                    "RunFullModel macro in gtperf.bas).").font = muted
    ws.freeze_panes = ws.cell(case_row + 1, first)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"B1:{get_column_letter(first + ncase - 1)}{row}"

    # Check sheet: full model at the default cases, stored as values, against the live formulas
    check_ws.sheet_view.showGridLines = False
    check_ws.column_dimensions["A"].width = 2
    check_ws.column_dimensions["B"].width = 26
    check_ws["B1"] = "Workbook formulas against the full gtperf model"
    check_ws["B1"].font = Font(bold=True, size=13, color=INK)
    check_ws["B2"] = ("Full-model values were computed when the workbook was generated, for the cases as "
                      "generated. If you change a case on the Site tool sheet the error column no longer applies.")
    check_ws["B2"].font = muted
    heads = ["Case", "Quantity", "Full model", "Workbook", "Error"]
    for j, h in enumerate(heads, start=2):
        c = check_ws.cell(4, j, h)
        c.font, c.fill = head, head_fill
        check_ws.column_dimensions[get_column_letter(j)].width = max(14, check_ws.column_dimensions[
            get_column_letter(j)].width or 0)
    r = 5
    units = {"power": ("Output, MW", "0.000", "0.00%"), "heat_rate": ("Heat rate, kJ/kWh", "#,##0.0", "0.00%"),
             "exhaust_flow": ("Exhaust flow, kg/s", "0.000", "0.00%"),
             "exhaust_temp": ("Exhaust temperature, °C", "0.0", '0.0 "K"')}
    for j, case in enumerate(cases):
        full = full_model(engine, case)
        for q in OUTPUTS:
            label, fmt, efmt = units[q]
            check_ws.cell(r, 2, case.name)
            check_ws.cell(r, 3, label)
            check_ws.cell(r, 4, full[q]).number_format = fmt
            check_ws.cell(r, 5, f"='Site tool'!{ref(res[q], j)}").number_format = fmt
            err = f"=E{r}-D{r}" if q == "exhaust_temp" else f"=E{r}/D{r}-1"
            check_ws.cell(r, 6, err).number_format = efmt
            r += 1
    check_ws.column_dimensions["C"].width = 26

    # About sheet
    about_ws.column_dimensions["A"].width = 110
    lines = [
        ("gtperf site performance tool", Font(bold=True, size=13, color=INK)),
        (f"Engine: {name}. {source}", None),
        ("", None),
        ("Method", bold),
        ("The Tables sheet holds corrections computed from the calibrated gtperf off-design model at sea level, "
         "every 5 °C from -20 to 50 °C:", None),
        ("  t_*: output, heat rate and exhaust flow as ratios to ISO, exhaust temperature as a change in K.", None),
        ("  rh_*, inlet_*, exhaust_*: linear sensitivities to humidity and to inlet and exhaust pressure loss.", None),
        ("Altitude is applied as the ISA pressure ratio delta (output and flow scale with delta, heat rate and "
         "exhaust temperature do not), which is exact for dry air by similarity.", None),
        ("Losses are divided by delta because their effect depends on the loss as a fraction of ambient pressure.",
         None),
        ("Effects are applied in a fixed order and multiplied, matching gtperf.site.site_performance.", None),
        ("", None),
        ("Limits", bold),
        ("Base load, natural gas as calibrated, new and clean unless degradation is entered. Fuel changes, "
         "part load and transient behaviour need the full model.", None),
        ("", None),
        ("Regenerating", bold),
        ("gtperf excel examples/sgt400_13mw.toml gtperf_site_tool.xlsx", Font(name="Consolas", color=INK)),
        ("", None),
        ("VBA module (optional)", bold),
        ("excel/gtperf.bas adds worksheet functions (GT_INTERP, GT_ISA_PRESSURE, GT_SITE_POWER and others) and a "
         "RunFullModel macro that sends each case to the Python model and writes the results under the sheet. "
         "Import it with Alt+F11, File > Import File, then save as .xlsm.", None),
        ("", None),
        ("All engine data are from public datasheets or synthetic. Not for commercial use.", muted),
        ("", None),
        ("Settings for RunFullModel", bold),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = about_ws.cell(i, 1, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            c.font = font
    r = len(lines) + 1
    for label, value, nm in (("gtperf command", command, "GTPERF_COMMAND"),
                             ("Datasheet file", str(datasheet_path), "DATASHEET_PATH")):
        about_ws.cell(r, 1, label).font = Font(color=INK2)
        c = about_ws.cell(r + 1, 1, value)
        c.fill, c.border = input_fill, box
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"About!$A${r + 1}")
        r += 2
    wb.calculation.fullCalcOnLoad = True  # no cached values are written, so ask Excel to calculate on open
    wb.properties.title = f"gtperf site performance: {name}"
    wb.properties.creator = "gtperf"
    wb.save(path)
    return tab


def cases_to_dicts(cases):
    return [asdict(c) for c in cases]
